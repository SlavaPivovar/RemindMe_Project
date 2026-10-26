"""База клієнтів і розсилка в Telegram."""

import io

import httpx
import openpyxl
from fastapi.testclient import TestClient

from tesla_ro.clients import build_clients, personalize
from tesla_ro.telegram import TelegramAPI, TelegramBot
from tesla_ro.web import create_app

from test_send_client import FakeTelegram, private_start


def make_order(session, name, phone, vin, status="прийнято", plate=""):
    session.new_order()
    session.set_customer(name, phone)
    session.set_vehicle(vin, odometer_km=10000, plate=plate)
    session.add_complaint("ТО")
    session.add_operation("BRK-PAD-F")
    session.save_order(status)
    return session.order.number


def test_build_clients_groups_by_phone(session):
    make_order(session, "Петренко Петро", "+380 67 123 45 67", "7SAYGDEEXPF000001", plate="АА1111АА")
    make_order(session, "Петренко Петро Іванович", "067 123 45 67", "5YJ3E1EA7KF000002", status="видано")
    make_order(session, "Чернетка Без Телефону", "", "5YJ3E1EA7KF000003", status="чернетка")
    make_order(session, "Скасований", "050 000 00 01", "5YJ3E1EA7KF000004", status="скасовано")
    clients = build_clients(session.store.list(), {"671234567": {"chat_id": 7, "name": "Петро", "at": "01.01.2026 10:00"},
                                                   "999999999": {"chat_id": 8, "name": "Без замовлень"}})
    by_key = {c["key"]: c for c in clients}
    petro = by_key["671234567"]
    assert petro["name"] == "Петренко Петро Іванович"  # новіші дані перекривають
    assert petro["visits"] == 2 and petro["spent"] == "444.00" and len(petro["cars"]) == 2
    assert petro["telegram"]["name"] == "Петро" and len(petro["orders"]) == 2
    assert by_key["n:чернетка без телефону"]["visits"] == 0  # чернетка — не візит
    assert by_key["500000001"]["visits"] == 0 and by_key["500000001"]["spent"] == "0.00"
    assert by_key["999999999"]["name"] == "Без замовлень" and by_key["999999999"]["orders"] == []
    text = personalize("Добрий день, {ім'я}! {авто} — {сервіс}, {адреса}", petro, {"name": "DK", "address": "Одеса"})
    assert text.startswith("Добрий день, Петренко Петро Іванович! Tesla Model") and text.endswith("— DK, Одеса")


def app_with_bot(tmp_path, session):
    fake = FakeTelegram()
    api = TelegramAPI("TOKEN", httpx.Client(transport=httpx.MockTransport(fake.handler)))
    app = create_app(orders_dir=session.store.dir,
                     telegram=lambda store, comments: TelegramBot(api, -100, store, comments, tmp_path / "tg.json"))
    return fake, TestClient(app), app.state.telegram


def test_clients_api_new_order_and_broadcast(tmp_path, session):
    number = make_order(session, "Петренко Петро", "+380 67 123 45 67", "7SAYGDEEXPF000001", plate="АА1111АА")
    make_order(session, "Іван Сидоренко", "+380 50 765 43 21", "5YJ3E1EA7KF000002")
    fake, web, bot = app_with_bot(tmp_path, session)
    token = web.post(f"/api/orders/{number}/client-link").json()["link"].split("start=")[1]
    bot.handle_update(private_start(token, 777))

    data = web.get("/api/clients").json()
    assert data["telegram"] and len(data["clients"]) == 2
    petro = next(c for c in data["clients"] if c["key"] == "671234567")
    assert petro["telegram"]["at"] and petro["months"] is not None and petro["months"] < 1

    # нове замовлення для відомого клієнта з вибраним автомобілем
    order = web.post("/api/clients/new-order", json={"key": "671234567", "vin": "7SAYGDEEXPF000001"}).json()["order"]
    assert order["number"] != number and order["customer"]["phone"] == "+380 67 123 45 67"
    assert order["vehicle"]["vin"] == "7SAYGDEEXPF000001" and order["vehicle"]["plate"] == "АА1111АА"
    assert order["vehicle"]["odometer_km"] is None  # пробіг майстер уточнює
    assert web.post("/api/clients/new-order", json={"key": "000"}).status_code == 404

    # розсилка: лише підключеним
    assert web.post("/api/broadcast", json={"keys": ["671234567"], "text": "  "}).status_code == 400
    r = web.post("/api/broadcast", json={"keys": ["671234567", "507654321"], "text": "Добрий день, {ім'я}!"}).json()
    assert r["sent"] == 1 and r["skipped"] == 1
    sent = [m for m in fake.sent("sendMessage") if m["chat_id"] == 777 and m["text"].startswith("Добрий день")]
    assert sent[-1]["text"] == "Добрий день, Петренко Петро!"
    assert web.get("/api/clients").json()["broadcasts"][0]["sent"] == 1
    assert any("Розсилка" in m["text"] for m in fake.sent("sendMessage") if m["chat_id"] == -100)

    x = web.get("/api/clients/export.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(x.content))["Клієнти"]
    assert ws["A1"].value == "ПІБ" and ws.max_row == 3


def test_broadcast_marks_blocked(tmp_path, session):
    number = make_order(session, "Петренко Петро", "+380 67 123 45 67", "7SAYGDEEXPF000001")
    fake, web, bot = app_with_bot(tmp_path, session)
    token = web.post(f"/api/orders/{number}/client-link").json()["link"].split("start=")[1]
    bot.handle_update(private_start(token, 777))
    original = fake.handler

    def blocked(request):
        if request.url.path.endswith("sendMessage") and b'"chat_id":777' in request.content.replace(b" ", b""):
            return httpx.Response(403, json={"ok": False, "description": "Forbidden: bot was blocked by the user"})
        return original(request)

    bot.api.http = httpx.Client(transport=httpx.MockTransport(blocked))
    r = web.post("/api/broadcast", json={"keys": ["671234567"], "text": "Акція"}).json()
    assert r["sent"] == 0 and r["failed"] == 1
    assert web.get("/api/clients").json()["clients"][0]["telegram"]["blocked"] is True
    assert web.post("/api/broadcast", json={"keys": ["671234567"], "text": "Акція"}).json()["skipped"] == 1


def test_no_bot(session):
    make_order(session, "Петренко Петро", "+380 67 123 45 67", "7SAYGDEEXPF000001")
    web = TestClient(create_app(orders_dir=session.store.dir))
    data = web.get("/api/clients").json()
    assert not data["telegram"] and data["clients"][0]["telegram"] is None
    assert web.post("/api/broadcast", json={"keys": ["671234567"], "text": "x"}).status_code == 503
