"""Фото, погодження кошторису, «автомобіль готовий», резервна копія й прайс в Excel."""

import base64
import io
import json
import zipfile

import httpx
import openpyxl
from fastapi.testclient import TestClient

from tesla_ro.telegram import TelegramAPI, TelegramBot
from tesla_ro.web import create_app

from test_send_client import GROUP, FakeTelegram, private_start, saved

JPEG = b"\xff\xd8\xff\xe0" + b"0" * 200
DATA_URL = "data:image/jpeg;base64," + base64.b64encode(JPEG).decode()
OWNER = 4242


class Fake(FakeTelegram):
    def handler(self, request):
        if "/file/" in request.url.path:
            return httpx.Response(200, content=JPEG)
        method = request.url.path.rsplit("/", 1)[-1]
        if method == "getFile":
            self.calls.append((method, json.loads(request.content)))
            return httpx.Response(200, json={"ok": True, "result": {"file_path": "photos/file_1.jpg"}})
        if method in ("sendPhoto", "sendMediaGroup"):
            self.calls.append((method, {"multipart": request.content}))
            return httpx.Response(200, json={"ok": True, "result": [{"message_id": 1}]})
        return super().handler(request)


def make(tmp_path, session, group=GROUP, backup_chat=0):
    fake = Fake()
    api = TelegramAPI("TOKEN", httpx.Client(transport=httpx.MockTransport(fake.handler)))

    def factory(store, comments):
        bot = TelegramBot(api, group, store, comments, tmp_path / "tg.json")
        bot.backup_chat_id = backup_chat
        return bot

    app = create_app(orders_dir=session.store.dir, telegram=factory)
    return fake, TestClient(app), app.state.telegram


def connect_client(web, bot, number, chat_id=777):
    token = web.post(f"/api/orders/{number}/client-link").json()["link"].split("start=")[1]
    bot.handle_update(private_start(token, chat_id))


# --- фото ---

def test_photos_web_and_client(tmp_path, session):
    number = saved(session)
    fake, web, bot = make(tmp_path, session)
    assert web.post("/api/orders/ЗН-2026-09999/photos", json={"image": DATA_URL}).status_code == 404
    items = web.post(f"/api/orders/{number}/photos", json={"image": DATA_URL, "kind": "пошкодження",
                                                           "author": "Іван"}).json()
    assert len(items) == 1 and items[0]["kind"] == "пошкодження" and items[0]["to_client"] is False
    pid = items[0]["id"]
    assert web.get(f"/photos/{number}/{pid}").content == JPEG
    assert web.post(f"/api/orders/{number}/photos", json={"image": "data:text/plain;base64,AA=="}).status_code == 400
    assert web.patch(f"/api/orders/{number}/photos/{pid}", json={"kind": "щось"}).status_code == 400
    web.post(f"/api/orders/{number}/photos", json={"image": DATA_URL, "kind": "до"})
    items = web.patch(f"/api/orders/{number}/photos/{pid}", json={"to_client": True}).json()
    assert [x["to_client"] for x in items] == [True, False]
    # фото потрапляють і в стан панелі
    web.post(f"/api/orders/{number}/open")
    assert len(web.get("/api/state").json()["order"]["photos"]) == 2

    connect_client(web, bot, number)
    r = web.post(f"/api/orders/{number}/send-telegram").json()
    assert r["photos"] == 1 and len(fake.sent("sendPhoto")) == 1  # лише позначене «клієнту»

    assert len(web.delete(f"/api/orders/{number}/photos/{pid}").json()) == 1
    assert web.get(f"/photos/{number}/{pid}").status_code == 404


def test_photo_from_group(tmp_path, session):
    number = saved(session)
    fake, web, bot = make(tmp_path, session)
    bot.order_saved(session.store.load(number))
    bot.flush()
    card = bot.links["orders"][number]["message_id"]
    bot.handle_update({"update_id": 5, "message": {
        "message_id": 90, "chat": {"id": GROUP, "type": "supergroup"}, "from": {"first_name": "Іван"},
        "caption": "після", "reply_to_message": {"message_id": card},
        "photo": [{"file_id": "small", "width": 90, "height": 90}, {"file_id": "big", "width": 1280, "height": 960}]}})
    photos = web.get(f"/api/orders/{number}/extras").json()["photos"]
    assert len(photos) == 1 and photos[0]["kind"] == "після" and photos[0]["source"] == "telegram"
    assert fake.sent("getFile")[0]["file_id"] == "big"
    assert not web.get(f"/api/orders/{number}/extras").json()["comments"]  # підпис «після» — не коментар


# --- погодження кошторису ---

def callback(number, answer, chat_id=777, message_id=None, qid="q1"):
    return {"update_id": 9, "callback_query": {"id": qid, "data": f"ap:{number}:{answer}", "from": {"first_name": "Петро"},
            "message": {"message_id": message_id, "chat": {"id": chat_id, "type": "private"}}}}


def test_estimate_approval(tmp_path, session):
    number = saved(session)
    fake, web, bot = make(tmp_path, session)
    bot.order_saved(session.store.load(number))
    bot.flush()
    # клієнт ще не в боті — погодження все одно створюється, клієнту — посилання на сторінку
    first = web.post(f"/api/orders/{number}/send-approval").json()
    assert first["telegram"] is None and "/c/" in first["link"] and not fake.sent("sendDocument")
    connect_client(web, bot, number)
    r = web.post(f"/api/orders/{number}/send-approval", headers={"x-forwarded-proto": "https", "host": "sto.example"}).json()
    approval = r["approval"]
    assert r["telegram"] is True and r["link"].startswith("https://sto.example/c/")
    assert approval["status"] == "pending" and approval["total"] == "222.00" and len(approval["lines"]) == 2
    doc = fake.sent("sendDocument")[-1]["multipart"].decode("utf-8", "replace")
    assert "ap:" in doc and "inline_keyboard" in doc and "https://sto.example/c/" in doc  # кнопка «Обрати по пунктах»

    bot.handle_update(callback(number, "y", chat_id=999, message_id=approval["message_id"]))  # чужий чат
    assert bot.approval(number)["status"] == "pending"
    bot.handle_update(callback(number, "y", message_id=approval["message_id"]))
    assert bot.approval(number)["status"] == "approved"
    bot.handle_update(callback(number, "n", message_id=approval["message_id"], qid="q2"))  # вдруге — не змінює
    assert bot.approval(number)["status"] == "approved"
    answers = fake.sent("answerCallbackQuery")
    assert "вже відповіли" in answers[-1]["text"]
    extras = web.get(f"/api/orders/{number}/extras").json()
    assert extras["approval"]["status"] == "approved"
    assert "погодив кошторис" in extras["comments"][-1]["text"]
    assert any("погодив" in m["text"] for m in fake.sent("sendMessage") if m["chat_id"] == GROUP)


# --- «Ваш автомобіль готовий» ---

def test_ready_notification_once(tmp_path, session):
    number = saved(session)
    fake, web, bot = make(tmp_path, session, group=0)  # працює і без групи
    connect_client(web, bot, number)
    for status in ("в роботі", "готово", "готово"):
        session.save_order(status)
        bot.order_saved(session.order)  # у сеансу тесту своє сховище, без обробників
        bot.flush()
        if status == "в роботі":
            assert not [m for m in fake.sent("sendMessage") if "готовий" in m["text"]]
    ready = [m for m in fake.sent("sendMessage") if "готовий" in m["text"]]
    assert len(ready) == 1 and ready[0]["chat_id"] == 777 and "222,00 $" in ready[0]["text"]


# --- резервна копія ---

def test_backup_download_and_telegram(tmp_path, session):
    number = saved(session)
    fake, web, bot = make(tmp_path, session, backup_chat=OWNER)
    web.post(f"/api/orders/{number}/photos", json={"image": DATA_URL})
    web.post(f"/api/orders/{number}/comments", json={"text": "Готово", "author": "Іван"})
    r = web.get("/api/backup.zip")
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert f"orders/{number}.json" in names and f"orders/comments/{number}.jsonl" in names
    assert any(n.startswith(f"orders/photos/{number}/") for n in names)

    assert web.post("/api/backup/telegram").json()["orders"] == 1
    # нічна копія: один раз на добу
    bot.links["backup_date"] = ""
    bot._backup_due = lambda: bot.links["backup_date"] == ""
    bot.backup_tick()
    bot.backup_tick()
    assert len(fake.sent("sendDocument")) == 2

    # /backup у приватному чаті — лише власнику
    bot.handle_update({"update_id": 1, "message": {"message_id": 1, "chat": {"id": 5, "type": "private"},
                                                   "text": "/backup", "from": {"id": 5}}})
    assert "лише власнику" in fake.sent("sendMessage")[-1]["text"] and len(fake.sent("sendDocument")) == 2
    bot.handle_update({"update_id": 2, "message": {"message_id": 2, "chat": {"id": OWNER, "type": "private"},
                                                   "text": "/backup", "from": {"id": OWNER}}})
    assert len(fake.sent("sendDocument")) == 3


def test_backup_skips_photos_over_limit(tmp_path):
    from tesla_ro.backup import make_backup

    orders = tmp_path / "orders"
    (orders / "photos" / "ЗН-2026-00001").mkdir(parents=True)
    (orders / "ЗН-2026-00001.json").write_text("{}")
    (orders / "photos" / "ЗН-2026-00001" / "a.jpg").write_bytes(b"x" * 5000)
    data, info = make_backup(orders, [], max_bytes=1000)
    assert info["photos"] == 0 and info["photos_skipped"] == 1 and info["orders"] == 1
    data, info = make_backup(orders, [])
    assert info["photos"] == 1


# --- прайс в Excel ---

def test_excel_export_import(session):
    web = TestClient(create_app(orders_dir=session.store.dir))
    r = web.get("/api/catalog/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ops = wb["Роботи"]
    assert ops["A1"].value == "Код" and ops.max_row > 50
    first_code = ops["A2"].value
    ops["D2"] = 123.45
    ops.append(["NEW-WORK", "Нова робота", "Інше", 40, "Model 3, Y", "", ""])
    ops.append(["BAD", "Без ціни", "Інше", "дорого", "усі"])
    wb["Запчастини"].append([1044321.0, "Нова деталь", 12.5, "шт"])
    buf = io.BytesIO()
    wb.save(buf)

    preview = web.post("/api/catalog/import", content=buf.getvalue()).json()
    assert {x["code"]: x["action"] for x in preview["operations"]} == {first_code: "changed", "NEW-WORK": "new"}
    assert preview["parts"][0]["code"] == "1044321" and len(preview["errors"]) == 1
    assert preview["errors"][0]["row"] == ops.max_row and not preview["applied"]
    assert "NEW-WORK" not in [o["code"] for o in web.get("/api/catalog").json()["operations"]]

    done = web.post("/api/catalog/import?apply=true", content=buf.getvalue()).json()
    ops_after = {o["code"]: o for o in done["catalog"]["operations"]}
    assert ops_after[first_code]["price"] == 123.45 and ops_after["NEW-WORK"]["model_codes"] == ["MODEL_3", "MODEL_Y"]
    # повторний імпорт того самого файлу — без змін
    again = web.post("/api/catalog/import", content=buf.getvalue()).json()
    assert not again["operations"] and not again["parts"]
    assert web.post("/api/catalog/import", content=b"not excel").status_code == 400
