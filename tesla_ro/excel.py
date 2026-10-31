"""Вивантаження довідника (прайсу) в Excel і завантаження змін назад.

Файл має два аркуші — «Роботи» і «Запчастини». Під час завантаження рядки звіряються за кодом
роботи чи артикулом: нові додаються, змінені оновлюються, рядки, яких у файлі немає, лишаються
в довіднику без змін (імпорт нічого не видаляє). Спершу показується попередній перегляд,
застосовуються лише рядки без помилок.
"""

from __future__ import annotations

import io
import re
from dataclasses import asdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .catalog import MODEL_CODES, Catalog, CatalogError, Operation, Part

OPS_SHEET, PARTS_SHEET = "Роботи", "Запчастини"
OPS_COLUMNS = [("code", "Код", 16), ("name", "Назва", 50), ("category", "Розділ", 22), ("price", "Ціна, $", 10),
               ("models", "Моделі", 24), ("parts", "Запчастини (артикули)", 30), ("interval", "Інтервал", 30)]
PARTS_COLUMNS = [("part_no", "Артикул", 18), ("name", "Назва", 50), ("price", "Ціна, $", 10), ("unit", "Одиниця", 10)]
ALL_MODELS = "усі"
MAX_ROWS = 5000


class ExcelError(ValueError):
    """Файл не вдалося прочитати; текст показується користувачу."""


def _models_text(models: list[str]) -> str:
    if "*" in models:
        return ALL_MODELS
    return ", ".join(m.replace("MODEL_", "Model ").replace("CYBERTRUCK", "Cybertruck") for m in models)


def _parse_models(text: str) -> list[str]:
    text = str(text or "").strip()
    if text.lower() in (ALL_MODELS, "усі моделі", "все", "*", "all"):
        return ["*"]
    models = []
    for item in re.split(r"[,;/]+", text):
        item = re.sub(r"[\s_]+", "", item.strip().upper()).replace("MODEL", "")
        if not item:
            continue
        code = "CYBERTRUCK" if item.startswith("CYBER") else f"MODEL_{item}"
        if code not in MODEL_CODES:
            raise CatalogError(f"Невідома модель «{item}». Можливі: S, 3, X, Y, Cybertruck або «{ALL_MODELS}»")
        models.append(code)
    return models


def _sheet(wb: Workbook, title: str, columns, rows) -> None:
    ws = wb.create_sheet(title)
    ws.append([c[1] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF5")
    for i, (_, _, width) in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"


def export_catalog(catalog: Catalog) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, OPS_SHEET, OPS_COLUMNS, [
        [o.code, o.name, o.category, o.price, _models_text(o.models), ", ".join(o.parts), o.interval]
        for o in catalog.operations.values()])
    _sheet(wb, PARTS_SHEET, PARTS_COLUMNS, [[p.part_no, p.name, p.price, p.unit] for p in catalog.parts.values()])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(value):
    # Excel зберігає числові артикули як 1044321.0 — повертаємо їм цілий вигляд.
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _read_rows(wb, title: str, columns) -> list[tuple[int, dict]]:
    if title not in wb.sheetnames:
        return []
    ws = wb[title]
    header = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    index = {}
    for key, label, _ in columns:
        # Колонки шукаємо за назвою — їх можна переставляти; «Ціна, $» знайдеться і як «Ціна».
        short = label.split(",")[0].split("(")[0].strip().lower()
        pos = next((i for i, h in enumerate(header) if h == label.lower() or h.split(",")[0].split("(")[0].strip() == short), None)
        if pos is not None:
            index[key] = pos
    if "code" not in index and "part_no" not in index:
        raise ExcelError(f"Аркуш «{title}»: не знайдено колонку «{columns[0][1]}» у першому рядку")
    rows = []
    for n, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if n > MAX_ROWS:
            raise ExcelError(f"Аркуш «{title}»: забагато рядків (до {MAX_ROWS})")
        row = {key: _cell(values[pos]) if pos < len(values) else None for key, pos in index.items()}
        if all(v is None or str(v).strip() == "" for v in row.values()):
            continue  # порожній рядок
        rows.append((n, row))
    return rows


def _changes(old, new) -> list[str]:
    if old is None:
        return []
    a, b = asdict(old), asdict(new)
    labels = {"name": "назва", "price": "ціна", "category": "розділ", "models": "моделі", "parts": "запчастини",
              "interval": "інтервал", "unit": "одиниця"}
    return [labels.get(k, k) for k in b if a.get(k) != b[k]]


def analyze(catalog: Catalog, data: bytes) -> dict:
    """Розібрати файл і порівняти з довідником. Повертає попередній перегляд і готові до запису рядки."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # openpyxl на пошкодженому чи не тому файлі кидає різні винятки
        raise ExcelError("Не вдалося прочитати файл. Потрібен файл Excel (.xlsx), вивантажений із довідника.") from None
    if OPS_SHEET not in wb.sheetnames and PARTS_SHEET not in wb.sheetnames:
        raise ExcelError(f"У файлі немає аркушів «{OPS_SHEET}» і «{PARTS_SHEET}»")
    errors: list[dict] = []
    parts: dict[str, Part] = {}
    preview = {"operations": [], "parts": []}

    for n, row in _read_rows(wb, PARTS_SHEET, PARTS_COLUMNS):
        try:
            part = catalog.make_part(row)
            if part.part_no in parts:
                raise CatalogError(f"Артикул {part.part_no} повторюється у файлі")
        except CatalogError as exc:
            errors.append({"sheet": PARTS_SHEET, "row": n, "message": str(exc)})
            continue
        parts[part.part_no] = part
        old = catalog.parts.get(part.part_no)
        if old != part:
            preview["parts"].append({"code": part.part_no, "name": part.name, "price": part.price,
                                     "action": "new" if old is None else "changed", "changes": _changes(old, part)})

    known_parts = {**catalog.parts, **parts}
    operations: dict[str, Operation] = {}
    for n, row in _read_rows(wb, OPS_SHEET, OPS_COLUMNS):
        try:
            row["models"] = _parse_models(row.get("models"))
            row["parts"] = [p for p in re.split(r"[,;\s]+", str(row.get("parts") or "")) if p]
            op = catalog.make_operation(row, known_parts)
            if op.code in operations:
                raise CatalogError(f"Код {op.code} повторюється у файлі")
        except CatalogError as exc:
            errors.append({"sheet": OPS_SHEET, "row": n, "message": str(exc)})
            continue
        operations[op.code] = op
        old = catalog.operations.get(op.code)
        if old != op:
            preview["operations"].append({"code": op.code, "name": op.name, "price": op.price,
                                          "action": "new" if old is None else "changed", "changes": _changes(old, op)})
    wb.close()
    changed_ops = [operations[x["code"]] for x in preview["operations"]]
    changed_parts = [parts[x["code"]] for x in preview["parts"]]
    return {"preview": {**preview, "errors": errors}, "operations": changed_ops, "parts": changed_parts}


def import_catalog(catalog: Catalog, data: bytes, apply: bool) -> dict:
    result = analyze(catalog, data)
    preview = result["preview"]
    if apply and (result["operations"] or result["parts"]):
        catalog.merge_items(result["operations"], result["parts"])
    preview["applied"] = bool(apply)
    return preview


CLIENT_COLUMNS = [("ПІБ", 30), ("Телефон", 18), ("E-mail", 24), ("Автомобілі", 40), ("Візитів", 9),
                  ("Сума, $", 12), ("Останній візит", 17), ("Telegram", 22)]


def export_clients(clients: list[dict]) -> bytes:
    from .clients import car_label

    wb = Workbook()
    wb.remove(wb.active)
    rows = []
    for c in clients:
        cars = "; ".join(" ".join(x for x in (car_label(car), car.get("vin", "")) if x) for car in c["cars"])
        tg = c["telegram"]
        rows.append([c["name"], c["phone"], c["email"], cars, c["visits"], float(c["spent"]),
                     c["last_visit"] or c["last_order"],
                     ("заблокував бота" if tg["blocked"] else f"так, з {tg['at']}" if tg["at"] else "так") if tg else ""])
    _sheet(wb, "Клієнти", [(None, label, width) for label, width in CLIENT_COLUMNS], rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_to_order(rows: list[dict]) -> bytes:
    """Що замовити: запчастини, яких на складі не більше мінімуму або не вистачає під замовлення."""
    wb = Workbook()
    wb.remove(wb.active)
    columns = [(None, "Артикул", 18), (None, "Назва", 50), (None, "Залишок", 10), (None, "У резерві", 10),
               (None, "Мінімум", 10), (None, "Замовити", 10), (None, "Одиниця", 10)]
    data = []
    for r in rows:
        need = max(r["min"] * 2 - r["free"], r["min"] - r["free"], -r["free"], 1)  # поповнити до двох мінімумів
        data.append([r["part_no"], r["name"], r["qty"], r["reserved"], r["min"], round(need, 3), r["unit"]])
    _sheet(wb, "Замовити", columns, data)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_purchase(po: dict, service: dict) -> bytes:
    """Замовлення постачальнику — щоб надіслати файлом."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Замовлення"
    ws.append([f"Замовлення {po['number']} від {po['created_at'].split(' ')[0]}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Замовник: {service.get('name', '')}" + (f", {service['address']}" if service.get("address") else "")])
    if po.get("supplier"):
        ws.append([f"Постачальник: {po['supplier']}"])
    ws.append([])
    ws.append(["Артикул", "Назва", "Кількість", "Одиниця", "Ціна, $", "Сума, $"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF5")
    for x in po["lines"]:
        ws.append([x["part_no"], x["name"], x["qty"], x.get("unit", "шт"), x["cost"] or None,
                   round(x["qty"] * x["cost"], 2) if x["cost"] else None])
    ws.append(["", "Разом", "", "", "", float(po["total"])])
    ws[ws.max_row][5].font = Font(bold=True)
    if po.get("note"):
        ws.append([])
        ws.append([f"Примітка: {po['note']}"])
    for col, width in zip("ABCDEF", (18, 50, 11, 10, 11, 12)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
